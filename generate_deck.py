import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation in 16:9 Widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Extracted directly from Aurora index.css & design system)
COLOR_BG = RGBColor(7, 10, 19)           # #070a13 Dark Space
COLOR_CARD_BG = RGBColor(17, 24, 39)      # #111827 Dark Slate Card
COLOR_CARD_BORDER = RGBColor(30, 41, 59)  # #1e293b Card Border
COLOR_ACCENT_CYAN = RGBColor(56, 189, 248)# #38bdf8 Cyan
COLOR_ACCENT_GREEN = RGBColor(16, 185, 129)# #10b981 Emerald
COLOR_ACCENT_RED = RGBColor(239, 68, 68)  # #ef4444 Red / Critical
COLOR_ACCENT_AMBER = RGBColor(245, 158, 11)# #f59e0b Amber / Warning
COLOR_TEXT_PRIMARY = RGBColor(248, 250, 252) # #f8fafc White/Ice
COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8 Slate Muted
COLOR_TEXT_DIM = RGBColor(100, 116, 139)   # #64748b Dim Slate

FONT_TITLE = "Inter"
FONT_BODY = "Inter"

PRESENTATION_DIR = "/home/kavya-singla/.gemini/antigravity-ide/scratch/safesignal/presentation"

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def add_header(slide, title_text, category_tag="AURORA SAFETY SYSTEM"):
    # Category Tag Pill
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_tag.upper()
    p_tag.font.name = FONT_TITLE
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_ACCENT_CYAN

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_PRIMARY

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

blank_slide_layout = prs.slide_layouts[6]

# ==============================================================================
# SLIDE 1: Title Slide
# ==============================================================================
slide1 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide1)

# Subtle background decorative card
add_card(slide1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=RGBColor(11, 17, 32), border_color=RGBColor(30, 41, 59))

# Tag badge
tag_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(1.4), Inches(3.2), Inches(0.4))
tag_shape.fill.solid()
tag_shape.fill.fore_color.rgb = RGBColor(30, 41, 59)
tag_shape.line.fill.background()
tf_t = tag_shape.text_frame
p_t = tf_t.paragraphs[0]
p_t.text = "🛡️ COMPOSITE RISK GUARDIAN"
p_t.font.name = FONT_TITLE
p_t.font.size = Pt(11)
p_t.font.bold = True
p_t.font.color.rgb = COLOR_ACCENT_CYAN
p_t.alignment = PP_ALIGN.CENTER

# Main Title
title_box = slide1.shapes.add_textbox(Inches(1.3), Inches(2.0), Inches(10.7), Inches(1.2))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Aurora"
p.font.name = FONT_TITLE
p.font.size = Pt(56)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_PRIMARY

# Subtitle / Tagline
sub_box = slide1.shapes.add_textbox(Inches(1.3), Inches(3.3), Inches(10.5), Inches(1.2))
tf_sub = sub_box.text_frame
tf_sub.word_wrap = True
p_sub = tf_sub.paragraphs[0]
p_sub.text = "A personal safety system that acts before you have to ask."
p_sub.font.name = FONT_BODY
p_sub.font.size = Pt(22)
p_sub.font.color.rgb = COLOR_ACCENT_CYAN

p_desc = tf_sub.add_paragraph()
p_desc.text = "Continuous multi-signal behavioral telemetry • Anti-coercion check-in • 100% on-device privacy"
p_desc.font.name = FONT_BODY
p_desc.font.size = Pt(14)
p_desc.font.color.rgb = COLOR_TEXT_MUTED
p_desc.space_before = Pt(12)

# Footer info
footer_box = slide1.shapes.add_textbox(Inches(1.3), Inches(5.3), Inches(10.5), Inches(0.8))
tf_foot = footer_box.text_frame
p_foot = tf_foot.paragraphs[0]
p_foot.text = "Hackathon Submission 2026  |  Production Deployment & Architecture Deck"
p_foot.font.name = FONT_BODY
p_foot.font.size = Pt(13)
p_foot.font.color.rgb = COLOR_TEXT_DIM

# ==============================================================================
# SLIDE 2: The Problem
# ==============================================================================
slide2 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide2)
add_header(slide2, "The Critical Flaw of Reactive Safety Apps", "PROBLEM STATEMENT")

problems = [
    ("🚨 The Manual Action Fallacy", "Most apps assume the victim can recognize danger and act mid-crisis (unlock phone, open app, press button). Under real panic, restraint, or distraction, that assumption fails."),
    ("🤐 The Coercion Blindspot", "Attackers routinely force victims to unlock phones and reply \"I'm fine.\" Existing check-in systems treat any reassuring response as safety confirmation and stand down."),
    ("⏱️ Hesitation & Fear of Embarrassment", "Users hesitate to trigger panic buttons out of social embarrassment or false-alarm fear until it is already too late for early prevention."),
    ("🔋 Always-On Fatigue & Disablement", "Constant 24/7 family GPS tracking causes heavy battery drain and surveillance discomfort on both ends, leading users to disable safety features entirely.")
]

for i, (p_title, p_desc) in enumerate(problems):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 5.95)
    top = Inches(1.5 + row * 2.65)
    
    add_card(slide2, left, top, Inches(5.75), Inches(2.45))
    
    tb = slide2.shapes.add_textbox(left + Inches(0.3), top + Inches(0.25), Inches(5.15), Inches(1.95))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = p_title
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ACCENT_AMBER
    
    p2 = tf.add_paragraph()
    p2.text = p_desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(8)

# ==============================================================================
# SLIDE 3: Our Approach
# ==============================================================================
slide3 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide3)
add_header(slide3, "Inverting the Safety Model: Default Toward Caution", "CORE PHILOSOPHY")

# Left Box: Traditional
add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.75), Inches(5.2), border_color=COLOR_CARD_BORDER)
tb_trad = slide3.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.15), Inches(4.7))
tf_trad = tb_trad.text_frame
tf_trad.word_wrap = True

p_t1 = tf_trad.paragraphs[0]
p_t1.text = "TRADITIONAL MODEL: REACTIVE"
p_t1.font.name = FONT_TITLE
p_t1.font.size = Pt(14)
p_t1.font.bold = True
p_t1.font.color.rgb = COLOR_ACCENT_RED

p_t2 = tf_trad.add_paragraph()
p_t2.text = "Assumes Safe Until Explicit Panic"
p_t2.font.name = FONT_TITLE
p_t2.font.size = Pt(20)
p_t2.font.bold = True
p_t2.font.color.rgb = COLOR_TEXT_PRIMARY
p_t2.space_before = Pt(6)

trad_points = [
    "• Passive state: App sleeps until manual trigger",
    "• Danger response: Waits for user to unlock and press SOS",
    "• Verification: Any response (\"I'm fine\") dismisses alert",
    "• Result: Fails completely if user is restrained or coerced"
]
for pt in trad_points:
    p = tf_trad.add_paragraph()
    p.text = pt
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(10)

# Right Box: Aurora
add_card(slide3, Inches(6.78), Inches(1.6), Inches(5.75), Inches(5.2), border_color=COLOR_ACCENT_CYAN)
tb_aurora = slide3.shapes.add_textbox(Inches(7.08), Inches(1.85), Inches(5.15), Inches(4.7))
tf_aurora = tb_aurora.text_frame
tf_aurora.word_wrap = True

p_a1 = tf_aurora.paragraphs[0]
p_a1.text = "AURORA MODEL: PROACTIVE"
p_a1.font.name = FONT_TITLE
p_a1.font.size = Pt(14)
p_a1.font.bold = True
p_a1.font.color.rgb = COLOR_ACCENT_GREEN

p_a2 = tf_aurora.add_paragraph()
p_a2.text = "Assumes Risk Upon Anomalies; Proof Required"
p_a2.font.name = FONT_TITLE
p_a2.font.size = Pt(20)
p_a2.font.bold = True
p_a2.font.color.rgb = COLOR_TEXT_PRIMARY
p_a2.space_before = Pt(6)

aurora_points = [
    "• Continuous telemetry: Evaluates speed, stillness, time, zones, battery",
    "• Multi-factor compounding: Correlated risks trigger quiet check-in",
    "• Strict safe-word stand-down: Only exact private word cancels alert",
    "• Fail-Closed: Silence or wrong word instantly escalates to SOS"
]
for pt in aurora_points:
    p = tf_aurora.add_paragraph()
    p.text = pt
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(10)

# ==============================================================================
# SLIDE 4: How the Risk Engine Works
# ==============================================================================
slide4 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide4)
add_header(slide4, "Composite Risk Scoring Engine (0 – 100)", "ENGINE ARCHITECTURE")

# 3 Columns
# Col 1: Signals Input
add_card(slide4, Inches(0.8), Inches(1.5), Inches(3.7), Inches(5.3))
tb1 = slide4.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(3.3), Inches(4.9))
tf1 = tb1.text_frame
tf1.word_wrap = True
tf1.paragraphs[0].text = "1. Passive Input Signals"
tf1.paragraphs[0].font.name = FONT_TITLE
tf1.paragraphs[0].font.size = Pt(16)
tf1.paragraphs[0].font.bold = True
tf1.paragraphs[0].font.color.rgb = COLOR_ACCENT_CYAN

signals_list = [
    "• Transit Pace (Walk / Accelerating / Sprint)",
    "• Stillness Duration (Stationary seconds)",
    "• Time of Day (Day vs Night Diurnal Cycle)",
    "• Route Familiarity (500m Geofenced Zones)",
    "• Battery Discharge (Gradual vs Instant Plunge)",
    "• Crowd Density (Isolated vs Populated Area)"
]
for s in signals_list:
    p = tf1.add_paragraph()
    p.text = s
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

# Col 2: Compounding Logic
add_card(slide4, Inches(4.75), Inches(1.5), Inches(3.7), Inches(5.3))
tb2 = slide4.shapes.add_textbox(Inches(4.95), Inches(1.7), Inches(3.3), Inches(4.9))
tf2 = tb2.text_frame
tf2.word_wrap = True
tf2.paragraphs[0].text = "2. Compounding & Decay"
tf2.paragraphs[0].font.name = FONT_TITLE
tf2.paragraphs[0].font.size = Pt(16)
tf2.paragraphs[0].font.bold = True
tf2.paragraphs[0].font.color.rgb = COLOR_ACCENT_AMBER

points_engine = [
    "• Correlated Cross-Signal Multipliers:",
    "   2 Factors: 1.15x  |  3 Factors: 1.25x",
    "   4+ Factors: 1.35x compounding multiplier",
    "• Asymmetric Response Dynamics:",
    "   Instant Escalation: Threat bursts jump score immediately without sluggish lag",
    "   Gradual Temporal Decay: 0.5 pts/sec (30 pts/min) de-escalation once calm returns"
]
for pe in points_engine:
    p = tf2.add_paragraph()
    p.text = pe
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

# Col 3: Worked Benchmark Scenario
add_card(slide4, Inches(8.7), Inches(1.5), Inches(3.8), Inches(5.3), border_color=COLOR_ACCENT_RED)
tb3 = slide4.shapes.add_textbox(Inches(8.9), Inches(1.7), Inches(3.4), Inches(4.9))
tf3 = tb3.text_frame
tf3.word_wrap = True
tf3.paragraphs[0].text = "3. Worked Threat Scenario"
tf3.paragraphs[0].font.name = FONT_TITLE
tf3.paragraphs[0].font.size = Pt(16)
tf3.paragraphs[0].font.bold = True
tf3.paragraphs[0].font.color.rgb = COLOR_ACCENT_RED

scenario_lines = [
    "Scenario: Night Evasive Sprint",
    "• Night Vulnerability Baseline (+15 pts)",
    "• Unfamiliar Route Perimeter (+12 pts)",
    "• Deserted Isolated Area (+10 pts)",
    "• Sudden Accelerating Pace (+15 pts)",
    "• 4-Factor Compounding (1.35x Multiplier)",
    "─────────────────────",
    "Formula: (15 + 12 + 10 + 15) x 1.35 = 70.2",
    "RESULT: Score 70 / 100  ->  [HIGH TIER]",
    "Trigger: Quiet Safe-Word Check-In Prompt"
]
for sl in scenario_lines:
    p = tf3.add_paragraph()
    p.text = sl
    p.font.name = FONT_BODY
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_TEXT_PRIMARY if "RESULT" in sl or "Formula" in sl else COLOR_TEXT_MUTED
    p.space_before = Pt(4)

# ==============================================================================
# SLIDE 5: Anti-Coercion Check-In
# ==============================================================================
slide5 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide5)
add_header(slide5, "Fail-Closed Anti-Coercion Protocol", "VERIFICATION PIPELINE")

add_card(slide5, Inches(0.8), Inches(1.5), Inches(11.733), Inches(1.4), bg_color=RGBColor(20, 30, 50))
tb_top = slide5.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.1))
tf_top = tb_top.text_frame
tf_top.word_wrap = True
p_top1 = tf_top.paragraphs[0]
p_top1.text = "CORE INSIGHT: Coerced Confirmations Are Inherently Untrustworthy"
p_top1.font.name = FONT_TITLE
p_top1.font.size = Pt(16)
p_top1.font.bold = True
p_top1.font.color.rgb = COLOR_ACCENT_CYAN
p_top2 = tf_top.add_paragraph()
p_top2.text = "In an active confrontation, a victim can easily be intimidated into typing \"I am okay\" or tapping an \"I'm safe\" button. Traditional apps cancel alerts upon receiving reassurance — Aurora treats unauthenticated reassurance as an active coercion indicator."
p_top2.font.name = FONT_BODY
p_top2.font.size = Pt(12.5)
p_top2.font.color.rgb = COLOR_TEXT_MUTED
p_top2.space_before = Pt(4)

# 3 Step Boxes Below
steps = [
    ("1. Elevated Risk Trigger", "When composite score crosses into Elevated/High tier, Aurora suppresses alarms and presents a discreet verification prompt without alerting bystanders.", COLOR_ACCENT_CYAN),
    ("2. Strict Secret Safe Word", "Only the user's secret arbitrary phrase (e.g. \"AURORA7\") stands down the alert. Typing \"I'm fine\" or any wrong word immediately triggers Emergency SOS.", COLOR_ACCENT_AMBER),
    ("3. Silence Escalation Timeline", "Complete silence gives 3 timed retry windows (30s window + 10s gap). If unanswered after 3 attempts, Emergency SOS automatically broadcasts.", COLOR_ACCENT_RED)
]

for i, (st, sd, sc) in enumerate(steps):
    left = Inches(0.8 + i * 4.0)
    add_card(slide5, left, Inches(3.1), Inches(3.733), Inches(3.7))
    tb = slide5.shapes.add_textbox(left + Inches(0.25), Inches(3.3), Inches(3.233), Inches(3.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = st
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = sc
    
    p2 = tf.add_paragraph()
    p2.text = sd
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(10)

# ==============================================================================
# SLIDE 6: Sharing Model (Session-Scoped vs Always-On)
# ==============================================================================
slide6 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide6)
add_header(slide6, "Session-Scoped Sharing vs Always-On Surveillance", "PRIVACY & ENGAGEMENT")

sharing_cards = [
    ("👤 Solo Mode", "Zero Sharing by Default", "Designed for maximum privacy. All risk evaluation, sensors, and state remain 100% on-device. Nothing is ever broadcast to anyone except when the user or silence escalates to local Emergency SOS links.", COLOR_ACCENT_CYAN),
    ("👥 Connected Mode", "Emergency-Only Circle", "Configures 1 to 3 trusted contacts. Circle is never notified during routine safe commutes. Contacts are only alerted during genuine emergencies (Critical tier transitions or Emergency SOS dispatch).", COLOR_ACCENT_GREEN),
    ("📍 Live Journey Share", "Time-Boxed Active Walk", "User explicitly starts monitoring before a specific walk (e.g. night transit) and it auto-terminates on arrival. Sharing is strictly scoped to moments the user chooses, completely eliminating notification fatigue.", COLOR_ACCENT_AMBER)
]

for i, (st, ss, sd, sc) in enumerate(sharing_cards):
    left = Inches(0.8 + i * 4.0)
    add_card(slide6, left, Inches(1.6), Inches(3.733), Inches(5.2))
    tb = slide6.shapes.add_textbox(left + Inches(0.25), Inches(1.85), Inches(3.233), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = st
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = sc
    
    p2 = tf.add_paragraph()
    p2.text = ss
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_PRIMARY
    p2.space_before = Pt(4)
    
    p3 = tf.add_paragraph()
    p3.text = sd
    p3.font.name = FONT_BODY
    p3.font.size = Pt(12.5)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.space_before = Pt(12)

# ==============================================================================
# SLIDE 7: What Makes Aurora Different
# ==============================================================================
slide7 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide7)
add_header(slide7, "Genuinely Novel Differentiators vs Market Apps", "MARKET COMPARISON")

# Reference context box
add_card(slide7, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.8), bg_color=RGBColor(15, 23, 42))
tb_m = slide7.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(0.7))
tf_m = tb_m.text_frame
tf_m.paragraphs[0].text = "Reference Benchmark: Established safety apps (Life360, Noonlight, bSafe, Safetipin) rely on manual panic buttons or static pre-computed neighborhood scores. None combine live multi-signal behavioral fusion with fail-closed coercion defense."
tf_m.paragraphs[0].font.name = FONT_BODY
tf_m.paragraphs[0].font.size = Pt(11.5)
tf_m.paragraphs[0].font.color.rgb = COLOR_ACCENT_CYAN

diffs = [
    ("1. Live Multi-Signal Fusion", "Existing apps score an area once in advance; Aurora scores a dynamic live moment as it happens by continuously compounding kinematics, diurnal time, geofencing, and battery telemetry."),
    ("2. Fail-Closed Safe Word", "Existing panic buttons treat any tap as safe; Aurora defaults to security — silence and false reassurance fail closed and immediately arm Emergency SOS."),
    ("3. Session-Scoped Privacy", "Eliminates standing 24/7 location tracking. Users activate Live Journey Sharing only when needed, avoiding notification fatigue and privacy erosion."),
    ("4. Reusable Modular Core", "Pure-function scoring engine decoupled from UI. Can swap risk 'lenses' (e.g. medical falls, pregnancy, elder care) without rewriting app architecture.")
]

for i, (dt, dd) in enumerate(diffs):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 5.95)
    top = Inches(2.45 + row * 2.3)
    
    add_card(slide7, left, top, Inches(5.75), Inches(2.15))
    tb = slide7.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.25), Inches(1.75))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = dt
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_PRIMARY
    
    p2 = tf.add_paragraph()
    p2.text = dd
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(6)

# ==============================================================================
# SLIDE 8: Live Product Screenshots
# ==============================================================================
slide8 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide8)
add_header(slide8, "Live Application Interfaces (Deployed & Running)", "PRODUCT DEMONSTRATION")

screens = [
    ("screen_risk_baseline.png", "1. Threat Radar & Guard HUD", "Circular gauge, diurnal & GPS chips, and AI contributing factors breakdown."),
    ("screen_checkin_elevated.png", "2. Anti-Coercion Check-In", "Discreet countdown prompt triggered upon elevated risk anomalies."),
    ("screen_sos_modal.png", "3. Emergency SOS & Direct Dispatch", "Tap-to-call links (112, 100) and strict safe-word disarm authentication.")
]

for i, (img_name, cap_title, cap_desc) in enumerate(screens):
    left = Inches(0.8 + i * 4.0)
    
    # Image container
    img_path = os.path.join(PRESENTATION_DIR, img_name)
    if os.path.exists(img_path):
        slide8.shapes.add_picture(img_path, left, Inches(1.5), width=Inches(3.733))
    else:
        add_card(slide8, left, Inches(1.5), Inches(3.733), Inches(4.3))
    
    # Caption box
    tb = slide8.shapes.add_textbox(left, Inches(5.9), Inches(3.733), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = cap_title
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ACCENT_CYAN
    
    p2 = tf.add_paragraph()
    p2.text = cap_desc
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(3)

# ==============================================================================
# SLIDE 9: Architecture
# ==============================================================================
slide9 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide9)
add_header(slide9, "Architecture: 100% On-Device & Zero Backend Attack Surface", "TECHNICAL DESIGN")

arch_blocks = [
    ("📡 Real Browser Telemetry", "• System Clock API (Diurnal Day/Night)\n• Geolocation API (Haversine Velocity, Stillness)\n• Battery Status API (Level & Anomaly Gradient)\n• Graceful Manual Fallback Dials if restricted", COLOR_ACCENT_CYAN),
    ("⚙️ Pure Function Risk Engine", "• Zero DOM, Zero React State, Zero External API\n• Input: Signals Object -> Output: Score & Tier\n• Fully reusable across Edge Workers, React Native, Wearable OS, and backend microservices", COLOR_ACCENT_AMBER),
    ("🔒 Local On-Device Persistence", "• All safe words, contact profiles, and zones persist strictly in browser localStorage\n• Zero database tracking, zero user tracking\n• Eliminates backend credential leak vulnerabilities", COLOR_ACCENT_GREEN)
]

for i, (at, ad, ac) in enumerate(arch_blocks):
    left = Inches(0.8 + i * 4.0)
    add_card(slide9, left, Inches(1.6), Inches(3.733), Inches(5.2))
    tb = slide9.shapes.add_textbox(left + Inches(0.25), Inches(1.85), Inches(3.233), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = at
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = ac
    
    p2 = tf.add_paragraph()
    p2.text = ad
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12.5)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(12)

# ==============================================================================
# SLIDE 10: Known Limitations
# ==============================================================================
slide10 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide10)
add_header(slide10, "Intellectual Honesty: Known Proof-of-Concept Constraints", "LIMITATIONS")

limits = [
    ("📡 Simulated Notification Delivery", "Trusted Circle notifications are currently simulated in the client UI. A production rollout requires dedicated SMS/Push gateways (e.g. Twilio, FCM) for off-device recipient alerting."),
    ("🚔 Direct Dispatch Integration", "Emergency flow surfaces native tap-to-call links (112, 100 in India). Automated PSAP dispatch integration requires governmental and public-safety institutional partnerships."),
    ("🗺️ Area Crime Statistics", "Route familiarity is currently derived from user-saved geofenced zones. Integrating macro historical incident crime data (e.g. NCRB records) remains future scope."),
    ("🛡️ Physical Coercion Ceiling", "Software cannot physically stop a determined in-person assailant. Safe-word protocols significantly raise the barrier to spoofing, but have physical real-world limits.")
]

for i, (lt, ld) in enumerate(limits):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 5.95)
    top = Inches(1.5 + row * 2.65)
    
    add_card(slide10, left, top, Inches(5.75), Inches(2.45))
    tb = slide10.shapes.add_textbox(left + Inches(0.3), top + Inches(0.25), Inches(5.15), Inches(1.95))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = lt
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ACCENT_AMBER
    
    p2 = tf.add_paragraph()
    p2.text = ld
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12.5)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(8)

# ==============================================================================
# SLIDE 11: Future Scope
# ==============================================================================
slide11 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide11)
add_header(slide11, "Engineering Roadmap & Production Extensions", "FUTURE SCOPE")

futures = [
    ("📊 Predictive NCRB Crime Heatmaps", "Ingesting historical police precinct records to automatically adjust geographical baseline risk without manual user geofencing."),
    ("🎙️ Ambient Distress Sound AI", "On-device lightweight audio classification (Google YAMNet) detecting screams or breaking glass as an emergency signal without storing audio."),
    ("👆 Duress-Specific Biometrics", "Alternate fingerprint or disguised spoken phrases that visually appear to confirm safety to an attacker but silently trigger emergency escalation."),
    ("⌚ Wearable Sensor Integration", "Smartwatch biometric monitoring (heart-rate spikes, sudden fall acceleration) and discreet physical confirmation tap gestures."),
    ("💬 Dedicated Messaging Gateway", "End-to-end encrypted SMS & automated voice dispatch through carrier gateways to deliver reliable alerts in low-connectivity areas.")
]

for i, (ft, fd) in enumerate(futures):
    top = Inches(1.5 + i * 1.08)
    add_card(slide11, Inches(0.8), top, Inches(11.733), Inches(0.95))
    tb = slide11.shapes.add_textbox(Inches(1.05), top + Inches(0.12), Inches(11.2), Inches(0.75))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = ft
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ACCENT_CYAN
    
    p2 = tf.add_paragraph()
    p2.text = fd
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(2)

# ==============================================================================
# SLIDE 12: Closing & Links
# ==============================================================================
slide12 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide12)

add_card(slide12, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=RGBColor(11, 17, 32), border_color=COLOR_ACCENT_CYAN)

tb_c = slide12.shapes.add_textbox(Inches(1.3), Inches(1.2), Inches(10.7), Inches(1.0))
tf_c = tb_c.text_frame
tf_c.word_wrap = True
p_c1 = tf_c.paragraphs[0]
p_c1.text = "Aurora: Safety Built for Reality"
p_c1.font.name = FONT_TITLE
p_c1.font.size = Pt(40)
p_c1.font.bold = True
p_c1.font.color.rgb = COLOR_TEXT_PRIMARY

p_c2 = tf_c.add_paragraph()
p_c2.text = "Because in a real crisis, you should never have to remember to ask for help."
p_c2.font.name = FONT_BODY
p_c2.font.size = Pt(18)
p_c2.font.color.rgb = COLOR_ACCENT_CYAN
p_c2.space_before = Pt(8)

# Links Box
add_card(slide12, Inches(1.3), Inches(2.9), Inches(10.7), Inches(2.0), bg_color=RGBColor(17, 24, 39), border_color=COLOR_CARD_BORDER)
tb_l = slide12.shapes.add_textbox(Inches(1.6), Inches(3.1), Inches(10.1), Inches(1.6))
tf_l = tb_l.text_frame
tf_l.word_wrap = True

p_l1 = tf_l.paragraphs[0]
p_l1.text = "🌐 LIVE PRODUCTION APP:"
p_l1.font.name = FONT_TITLE
p_l1.font.size = Pt(13)
p_l1.font.bold = True
p_l1.font.color.rgb = COLOR_ACCENT_GREEN

p_l1_url = tf_l.add_paragraph()
p_l1_url.text = "https://safesignal-five.vercel.app"
p_l1_url.font.name = FONT_BODY
p_l1_url.font.size = Pt(16)
p_l1_url.font.bold = True
p_l1_url.font.color.rgb = COLOR_TEXT_PRIMARY
p_l1_url.space_before = Pt(2)

p_l2 = tf_l.add_paragraph()
p_l2.text = "📦 GITHUB REPOSITORY (PUBLIC):"
p_l2.font.name = FONT_TITLE
p_l2.font.size = Pt(13)
p_l2.font.bold = True
p_l2.font.color.rgb = COLOR_ACCENT_CYAN
p_l2.space_before = Pt(10)

p_l2_url = tf_l.add_paragraph()
p_l2_url.text = "https://github.com/KavyaSingla244/Aurora"
p_l2_url.font.name = FONT_BODY
p_l2_url.font.size = Pt(16)
p_l2_url.font.bold = True
p_l2_url.font.color.rgb = COLOR_TEXT_PRIMARY
p_l2_url.space_before = Pt(2)

# Footnote
tb_cf = slide12.shapes.add_textbox(Inches(1.3), Inches(5.2), Inches(10.7), Inches(0.5))
tf_cf = tb_cf.text_frame
p_cf = tf_cf.paragraphs[0]
p_cf.text = "Ready for immediate evaluation • Includes interactive Judge Scenario Preset injector & live GPS feeds."
p_cf.font.name = FONT_BODY
p_cf.font.size = Pt(13)
p_cf.font.color.rgb = COLOR_TEXT_MUTED

# Save PPTX
output_path = os.path.join(PRESENTATION_DIR, "Aurora_Presentation_Deck.pptx")
prs.save(output_path)
print(f"✓ Generated PowerPoint presentation deck successfully at: {output_path}")
